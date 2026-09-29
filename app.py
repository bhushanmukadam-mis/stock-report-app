import io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import streamlit as st

st.set_page_config(page_title="Stock Report Generator", page_icon="📊")

st.title("📊 Anusaya Stock Report Generator")
st.write("Upload your raw inventory Excel file to generate the formatted stock report with subtotals and ageing.")

uploaded_file = st.file_uploader("Upload 'Anusaya Stock' Excel File", type=["xlsx", "xls"])

def process_excel(file_bytes):
    # 1. Read Sheet1 starting from Row 4 (skipping raw headers)
    df_raw = pd.read_excel(file_bytes, sheet_name='Sheet1', skiprows=3)

    # Clean column headers
    df_raw.columns = [str(col).strip().replace('\n', ' ') for col in df_raw.columns]

    # Target column mapping (supports both 'Stock Qty.' and 'Available Stock')
    target_cols = {
        'Inward Date': 'Inward Date',
        'Item Name': 'Item Name',
        'Marka': 'Marka',
        'Lot / Vakal': 'Lot / Vakal',
        'Count': 'Count',
        'Available Stock': 'Available Stock',
        'Stock Qty.': 'Available Stock'
    }

    col_map = {}
    for col in df_raw.columns:
        clean_col = col.strip()
        if clean_col in target_cols:
            col_map[col] = target_cols[clean_col]

    df = df_raw[list(col_map.keys())].rename(columns=col_map).copy()
    df = df.dropna(subset=['Inward Date', 'Item Name']).copy()

    # 2. Date conversion & Ageing calculation
    df['Inward Date DT'] = pd.to_datetime(df['Inward Date'], dayfirst=True)
    target_date = datetime(2026, 9, 10)
    df['Ageing'] = (target_date - df['Inward Date DT']).dt.days
    df['Inward Date Str'] = df['Inward Date DT'].dt.strftime('%d/%m/%Y')

    # Sort data sequentially
    df = df.sort_values(by=['Inward Date DT', 'Item Name'], kind='stable')

    # 3. Create Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_title = Font(name='Calibri', size=11, bold=True)
    font_header = Font(name='Calibri', size=11, bold=True)
    font_data = Font(name='Calibri', size=11, bold=False)
    font_total = Font(name='Calibri', size=11, bold=True)

    yellow_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    grey_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    thin_side = Side(border_style='thin', color='000000')

    align_top_left = Alignment(horizontal='left', vertical='top')
    align_center_vert = Alignment(horizontal='center', vertical='center')
    align_right_center = Alignment(horizontal='right', vertical='center')
    align_left_center = Alignment(horizontal='left', vertical='center')

    # Title Block (Row 1)
    ws.merge_cells("B1:G1")
    ws["B1"].value = "Anusaya Stock 10/09/2026"
    ws["B1"].font = font_title
    ws["B1"].alignment = align_center_vert

    # Table Headers (Row 2)
    headers = ["Inward Date", "Item Name", "Marka", "Lot / Vakal", "Count", "Available Stock", "Ageing"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = yellow_fill
        cell.alignment = align_top_left
        cell.border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)

    # 4. Populate Data Rows
    curr_row = 3

    for (inward_date, item_name), group in df.groupby(['Inward Date Str', 'Item Name'], sort=False):
        group_start_row = curr_row
        marka_val = group['Marka'].iloc[0] if pd.notna(group['Marka'].iloc[0]) else ""
        
        for _, row in group.iterrows():
            ws.cell(row=curr_row, column=1, value=row['Inward Date Str'])
            ws.cell(row=curr_row, column=2, value=row['Item Name'])
            ws.cell(row=curr_row, column=3, value=marka_val)
            
            ws.cell(row=curr_row, column=4, value=str(row['Lot / Vakal']) if pd.notna(row['Lot / Vakal']) else "").font = font_data
            ws.cell(row=curr_row, column=5, value=row['Count'] if pd.notna(row['Count']) else 0).font = font_data
            ws.cell(row=curr_row, column=5).alignment = align_right_center
            
            ws.cell(row=curr_row, column=6, value=row['Available Stock'] if pd.notna(row['Available Stock']) else 0).font = font_data
            ws.cell(row=curr_row, column=6).alignment = align_right_center
            
            ws.cell(row=curr_row, column=7, value=row['Ageing']).font = font_data
            ws.cell(row=curr_row, column=7).alignment = align_right_center

            curr_row += 1

        data_end_row = curr_row - 1

        # Merge Columns A, B, C vertically for group data rows
        if data_end_row > group_start_row:
            ws.merge_cells(start_row=group_start_row, start_column=1, end_row=data_end_row, end_column=1)
            ws.merge_cells(start_row=group_start_row, start_column=2, end_row=data_end_row, end_column=2)
            ws.merge_cells(start_row=group_start_row, start_column=3, end_row=data_end_row, end_column=3)

        for r in range(group_start_row, data_end_row + 1):
            for c in range(1, 4):
                ws.cell(row=r, column=c).font = font_data
                ws.cell(row=r, column=c).alignment = align_top_left

        # Set Borders for Data Rows
        for r in range(group_start_row, data_end_row + 1):
            for c in range(1, 8):
                cell = ws.cell(row=r, column=c)
                top_b = thin_side if (r == group_start_row or c in [4, 5, 6, 7]) else None
                bottom_b = thin_side if (r == data_end_row or c in [4, 5, 6, 7]) else None
                left_b = thin_side if c in [1, 2, 3, 4, 5, 6, 7] else None
                right_b = thin_side if c in [1, 2, 3, 4, 5, 6, 7] else None
                cell.border = Border(top=top_b, bottom=bottom_b, left=left_b, right=right_b)

        # 5. Insert Grey Highlighted Subtotal Row
        subtotal_row = curr_row
        
        tot_label = ws.cell(row=subtotal_row, column=2, value=f"{item_name} Total")
        tot_label.font = font_total
        tot_label.alignment = align_left_center
        
        tot_stock = ws.cell(row=subtotal_row, column=6, value=group['Available Stock'].sum())
        tot_stock.font = font_total
        tot_stock.alignment = align_right_center
        
        tot_age = ws.cell(row=subtotal_row, column=7, value=group['Ageing'].iloc[0])
        tot_age.font = font_total
        tot_age.alignment = align_right_center

        # Apply Grey Fill & Full Borders across all columns (A-G) in Total Row
        for c in range(1, 8):
            cell = ws.cell(row=subtotal_row, column=c)
            cell.fill = grey_fill
            cell.border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)

        curr_row += 1

    # 6. Insert Yellow Highlighted Grand Total Row
    grand_total_row = curr_row

    grand_label = ws.cell(row=grand_total_row, column=1, value="Grand Total")
    grand_label.font = font_total
    grand_label.alignment = align_left_center

    total_available_stock = df['Available Stock'].sum()
    grand_stock_cell = ws.cell(row=grand_total_row, column=6, value=total_available_stock)
    grand_stock_cell.font = font_total
    grand_stock_cell.alignment = align_right_center

    for c in range(1, 8):
        cell = ws.cell(row=grand_total_row, column=c)
        cell.fill = yellow_fill
        cell.border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)

    # 7. Adjust Column Widths
    col_widths = {'A': 14, 'B': 30, 'C': 22, 'D': 16, 'E': 8, 'F': 16, 'G': 8}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Save output to buffer
    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    return output_buffer

if uploaded_file is not None:
    try:
        with st.spinner("Processing file..."):
            processed_data = process_excel(uploaded_file)
        
        st.success("Report generated successfully!")
        
        st.download_button(
            label="📥 Download Processed Stock Report",
            data=processed_data,
            file_name="Processed_Anusaya_Stock.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        st.error(f"Error processing file: {e}")
